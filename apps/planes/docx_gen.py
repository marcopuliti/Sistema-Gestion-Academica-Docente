import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MESES = [
    'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
    'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
]
DIAS = {1: 'Lunes', 2: 'Martes', 3: 'Miércoles', 4: 'Jueves', 5: 'Viernes'}

COLOR_AZUL     = '1A3A5C'
COLOR_AMARILLO = 'FFF3CD'
COLOR_GRIS     = 'F0F4F8'
COLOR_BLANCO   = 'FFFFFF'


def _set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)


def _write_cell(cell, texto, bold=False, italic=False, size=9, color=None, center=False):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(str(texto) if texto else '—')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    _cell_margins(cell)


def _p(doc_or_para, text, size=11, bold=False, italic=False, align=None):
    if hasattr(doc_or_para, 'add_paragraph'):
        para = doc_or_para.add_paragraph()
    else:
        para = doc_or_para
    run = para.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if align:
        para.alignment = align
    return para


def _diffs(t_dir, t_adm):
    if t_adm is None:
        return {k: True for k in ('presidente', 'vocal_1', 'vocal_2', 'dia', 'hora', 'modalidad')}
    return {
        'presidente': (t_dir.presidente_nombre != t_adm.presidente_nombre or
                       t_dir.presidente_dni != t_adm.presidente_dni),
        'vocal_1':    (t_dir.vocal_1_nombre != t_adm.vocal_1_nombre or
                       t_dir.vocal_1_dni != t_adm.vocal_1_dni),
        'vocal_2':    (t_dir.vocal_2_nombre != t_adm.vocal_2_nombre or
                       t_dir.vocal_2_dni != t_adm.vocal_2_dni),
        'dia':        t_dir.dia_semana != t_adm.dia_semana,
        'hora':       t_dir.hora != t_adm.hora,
        'modalidad':  t_dir.permite_libres != t_adm.permite_libres,
    }


_ANCHOS_TURNO_CM = (7.25, 1.95, 7.2, 2.15, 3.55, 3.6)
COLOR_RESALTADO = 'FFFF00'


def _bordes_celda(cell, **lados):
    """Grosor de borde por lado (en octavos de punto), p. ej. top=12."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for lado, sz in lados.items():
        el = borders.find(qn(f'w:{lado}'))
        if el is None:
            el = OxmlElement(f'w:{lado}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:color'), '000000')


def _celda_turno(cell, partes, size=8, center=True, fondo=None):
    """partes: lista de líneas; cada línea es una lista de (texto, {'bold':…, 'underline':…})."""
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(0)
    for i, linea in enumerate(partes):
        if i:
            para.add_run().add_break()
        for texto, fmt in linea:
            run = para.add_run(texto)
            run.font.name = 'Arial'
            run.font.size = Pt(size)
            run.font.bold = fmt.get('bold', False)
            run.font.underline = fmt.get('underline', False)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _cell_margins(cell, top=10, bottom=10, left=50, right=50)
    if fondo:
        _set_cell_bg(cell, fondo)


def _add_tabla_turno(doc, departamento, dia, bloques):
    from .pdf import DIAS as _DIAS, lineas_asignatura, titulo_turno

    filas = 2 + sum(len(b['miembros']) for b in bloques)
    table = doc.add_table(rows=filas, cols=6)
    table.style = 'Table Grid'
    table.alignment = 1
    for row in table.rows:
        for cell, ancho in zip(row.cells, _ANCHOS_TURNO_CM):
            cell.width = Cm(ancho)

    _celda_turno(table.cell(0, 0).merge(table.cell(0, 5)), [[(titulo_turno(departamento, dia), {'bold': True})]], size=13)
    _celda_turno(table.cell(1, 0), [[('ASIGNATURA', {'bold': True})]])
    _celda_turno(table.cell(1, 1).merge(table.cell(1, 2)), [[('MESA EXAMINADORA', {'bold': True})]])
    _celda_turno(table.cell(1, 3), [[('DIA Y', {'bold': True})], [('HORA', {'bold': True})]])
    _celda_turno(table.cell(1, 4), [[('LIBRES', {})]])
    _celda_turno(table.cell(1, 5), [[('PLAN DE', {})], [('ESTUDIOS', {})]])

    r = 2
    for b in bloques:
        n = len(b['miembros'])
        ultima = r + n - 1
        asignatura = [
            [(texto, {'bold': True}), (resto, {})] if negrita else [(texto, {})]
            for negrita, texto, resto in lineas_asignatura(b)
        ]
        celda = table.cell(r, 0).merge(table.cell(ultima, 0))
        _celda_turno(celda, asignatura, center=False)
        celda.vertical_alignment = WD_ALIGN_VERTICAL.TOP

        if b['dia']:
            dia_hora = [[(_DIAS.get(b['dia'], ''), {'underline': True})]]
            if b['hora']:
                dia_hora += [[('', {})], [(f"{b['hora']:%H:%M}hs", {})]]
        else:
            dia_hora = [[('—', {})]]
        _celda_turno(table.cell(r, 3).merge(table.cell(ultima, 3)), dia_hora,
                     fondo=COLOR_RESALTADO if b['cambio_dia'] else None)
        _celda_turno(table.cell(r, 4).merge(table.cell(ultima, 4)), [[(b['libres'], {})]],
                     fondo=COLOR_RESALTADO if b['cambio_libres'] else None)
        _celda_turno(table.cell(r, 5).merge(table.cell(ultima, 5)), [[(pl, {})] for pl in b['planes']])

        for i, (rol, integrante, cambiado) in enumerate(b['miembros']):
            fondo = COLOR_RESALTADO if cambiado else None
            _celda_turno(table.cell(r + i, 1), [[(rol, {})]], fondo=fondo)
            table.cell(r + i, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _celda_turno(table.cell(r + i, 2), [[(integrante, {})]], center=False, fondo=fondo)

        for col in range(6):
            _bordes_celda(table.cell(r, col), top=12)
        r += n

    # Borde exterior grueso a nivel tabla (los bordes de celda se pierden al combinar)
    tblPr = table._tbl.tblPr
    bordes = tblPr.find(qn('w:tblBorders'))
    if bordes is None:
        bordes = OxmlElement('w:tblBorders')
        tblPr.append(bordes)
    for lado, sz in (('top', 16), ('left', 16), ('bottom', 16), ('right', 16), ('insideH', 4), ('insideV', 4)):
        el = bordes.find(qn(f'w:{lado}'))
        if el is None:
            el = OxmlElement(f'w:{lado}')
            bordes.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:color'), '000000')

    # Que el bloque título + encabezado no quede solo al final de una página
    for row in table.rows[:2]:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = True


def _encabezado_turno_docx(section):
    """Escudo de la UNSL y 'SAN LUIS, MES AÑO' en el encabezado de las páginas de tablas."""
    import os
    from django.conf import settings

    section.header.is_linked_to_previous = False
    hoy = datetime.date.today()
    tabla = section.header.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
    izq, der = tabla.rows[0].cells
    izq.width, der.width = Cm(5), Cm(20)
    escudo = os.path.join(settings.BASE_DIR, 'static', 'escudo.gif')
    para = izq.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(escudo):
        para.add_run().add_picture(escudo, height=Cm(1.5))
    for linea in ('Universidad Nacional de San Luis', 'Facultad de Ciencias Físico', 'Matemáticas y Naturales'):
        p = izq.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linea)
        run.font.name = 'Arial'
        run.font.size = Pt(6.5)
    der.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = der.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'SAN LUIS, {MESES[hoy.month - 1].upper()} {hoy.year}')
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    section.header.paragraphs[0].text = ''


def generar_docx_solicitud_cambio(director, admin, meps, item_map, current_map):
    hoy = datetime.date.today()
    n = len([m for m in meps if item_map.get(m.id)])
    cuerpo = (
        f"Me dirijo a Ud., y por su intermedio a quien corresponda, para solicitar "
        f"la modificación de {n} Tribunal{'es' if n != 1 else ''} "
        f"Examinador{'es' if n != 1 else ''} para mesas de examen del Ciclo Lectivo "
        f"{hoy.year} de las materias del Departamento de {director.departamento}. "
        f"Las celdas resaltadas en amarillo indican los campos propuestos a modificar "
        f"respecto a los valores actuales."
    )
    return _docx_tribunales(director, admin, meps, item_map, current_map, cuerpo)


def generar_docx_informe_tribunales(director, admin, meps, tribunal_map):
    """Informe completo de tribunales del departamento (sin resaltados)."""
    hoy = datetime.date.today()
    cuerpo = (
        f"Me dirijo a Ud., y por su intermedio a quien corresponda, para elevar "
        f"el listado de Tribunales Examinadores para mesas de examen del Ciclo "
        f"Lectivo {hoy.year} de las materias del Departamento de "
        f"{director.departamento}, informando los cambios y agregados realizados "
        f"a los mismos."
    )
    return _docx_tribunales(director, admin, meps, tribunal_map, tribunal_map, cuerpo)


def _docx_tribunales(director, admin, meps, item_map, current_map, cuerpo):
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from .pdf import bloques_tribunales

    doc = Document()

    # Márgenes y A4, como el PDF
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    hoy = datetime.date.today()
    fecha_str = f"San Luis, {hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}"

    # ── Carta ──
    p_fecha = _p(doc, fecha_str, size=11)
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    admin_nombre = admin.get_full_name() if admin else 'Secretario Académico'
    _p(doc, f"Sr. Secretario Académico de la FCFMyN", bold=True)
    _p(doc, admin_nombre)
    _p(doc, "S / D:")
    doc.add_paragraph()

    p_cuerpo = _p(doc, cuerpo)
    p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()
    _p(doc, "Sin otro particular, saludo a Ud. con atenta consideración.")
    doc.add_paragraph()
    doc.add_paragraph()

    firma = (
        f"{director.get_full_name()}\n"
        f"Director del Departamento de {director.departamento}\n"
        f"F.C.F.M.yN."
    )
    _p(doc, firma)

    # ── Listado por día, en páginas apaisadas ──
    apaisada = doc.add_section(WD_SECTION.NEW_PAGE)
    apaisada.orientation = WD_ORIENT.LANDSCAPE
    apaisada.page_width, apaisada.page_height = Cm(29.7), Cm(21)
    apaisada.left_margin = apaisada.right_margin = Cm(1.8)
    apaisada.top_margin, apaisada.bottom_margin = Cm(1.2), Cm(1.2)
    apaisada.header_distance = Cm(0.8)
    _encabezado_turno_docx(apaisada)

    for i, (dia, bloques) in enumerate(bloques_tribunales(meps, item_map, current_map)):
        if i:
            doc.add_page_break()
        _add_tabla_turno(doc, director.departamento, dia, bloques)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ── Informe cuatrimestral de comisiones ──────────────────────────────────────
# Mismo formato que el PDF (pdf.datos_tabla_comision): encabezado "DEPARTAMENTO DE …",
# columnas MATERIA | COMISIÓN | PERIODO | CONDICIÓN | CRÉDITO HORARIO | PLAN DE ESTUDIOS.

_ANCHOS_COMISION_CM = (4.25, 2.2, 8.15, 2.0, 2.25, 2.0, 2.15)


def _celda_comision(cell, texto, size=7.5, bold=False, center=True):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(0)
    for i, linea in enumerate((texto or '').split('\n')):
        if i:
            para.add_run().add_break()
        run = para.add_run(linea)
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.bold = bold
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _cell_margins(cell, top=20, bottom=20, left=50, right=50)


def _add_comision_table(doc, c, departamento):
    from .pdf import datos_tabla_comision
    d = datos_tabla_comision(c)
    n = len(d['equipo'])

    table = doc.add_table(rows=2 + n, cols=7)
    table.style = 'Table Grid'
    table.alignment = 1  # centrada
    for row in table.rows:
        for cell, ancho in zip(row.cells, _ANCHOS_COMISION_CM):
            cell.width = Cm(ancho)

    _celda_comision(table.cell(0, 0).merge(table.cell(0, 6)), f'DEPARTAMENTO DE {departamento.upper()}', size=8, bold=True)
    _celda_comision(table.cell(1, 0), 'MATERIA', size=8, bold=True)
    _celda_comision(table.cell(1, 1).merge(table.cell(1, 2)), 'COMISIÓN', size=8, bold=True)
    for col, texto in ((3, 'PERIODO'), (4, 'CONDICIÓN'), (5, 'CRÉDITO\nHORARIO'), (6, 'PLAN DE\nESTUDIOS')):
        _celda_comision(table.cell(1, col), texto, size=7)

    for i, (materia, (rol, docente)) in enumerate(zip(d['materia'], d['equipo'])):
        r = 2 + i
        _celda_comision(table.cell(r, 0), materia, bold=(i == 2))
        _celda_comision(table.cell(r, 1), rol, size=8, center=False)
        _celda_comision(table.cell(r, 2), docente)
    ultima = 1 + n
    for col, clave in ((3, 'periodo'), (4, 'condicion'), (5, 'credito'), (6, 'plan')):
        _celda_comision(table.cell(2, col).merge(table.cell(ultima, col)), d[clave])

    # Que la tabla no se corte entre páginas
    for row in table.rows[:-1]:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = True

    doc.add_paragraph()


def generar_docx_informe_comisiones(departamento, director, admin, solicitud, comisiones):
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from .pdf import texto_periodo_informe

    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4, como el PDF
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)

    hoy = datetime.date.today()
    p_fecha = _p(doc, f"San Luis, {hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}", size=11)
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()

    admin_nombre = admin.get_full_name() if admin else 'Secretario Académico'
    _p(doc, "Sr. Secretario Académico de la FCFMyN", bold=True)
    _p(doc, admin_nombre)
    _p(doc, "S / D:")
    doc.add_paragraph()

    n = len(comisiones)
    p_cuerpo = _p(doc, (
        f"Me dirijo a Ud., y por su intermedio a quien corresponda, para elevar el informe de "
        f"comisiones del {texto_periodo_informe(solicitud)} de las materias del Departamento de "
        f"{departamento}, que comprende {n} {'comisiones' if n != 1 else 'comisión'} con su "
        f"condición de cursada, crédito horario, planes de estudio asociados y equipo docente."
    ))
    p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    _p(doc, "Sin otro particular, saludo a Ud. con atenta consideración.")
    doc.add_paragraph()
    doc.add_paragraph()
    nombre_director = director.get_full_name() if director else ''
    _p(doc, f"{nombre_director}\nDirector del Departamento de {departamento}\nF.C.F.M.yN.")

    # Tablas en páginas apaisadas
    apaisada = doc.add_section(WD_SECTION.NEW_PAGE)
    apaisada.orientation = WD_ORIENT.LANDSCAPE
    apaisada.page_width, apaisada.page_height = apaisada.page_height, apaisada.page_width
    apaisada.left_margin = apaisada.right_margin = Cm(2)
    apaisada.top_margin = apaisada.bottom_margin = Cm(1.5)
    for c in comisiones:
        _add_comision_table(doc, c, departamento)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
